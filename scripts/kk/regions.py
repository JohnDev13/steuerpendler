"""Convert the BAG lists of Gemeinden by premium region (data/kk/raw/praemienregionen-<year>.xlsx, annex of the EDI
ordinance SR 832.106) to data/kk/regions-<year>.csv (bfs,canton,name,region) and the merger table to
data/kk/mergers-<year>.csv (bfs,canton,name,into). Cantons without regions are not listed by BAG (region 0).

    python3 scripts/kk/regions.py 2026 2027        # needs openpyxl
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]

for year in sys.argv[1:]:
    wb = openpyxl.load_workbook(ROOT / f"data/kk/raw/praemienregionen-{year}.xlsx", read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[0] and str(r[0]).startswith("Kanton")) + 1
    out = [(int(r[1]), r[0], str(r[2]).strip(), int(r[3])) for r in rows[start:] if r[0] and r[1]]
    with open(ROOT / f"data/kk/regions-{year}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n"); w.writerow(["bfs", "canton", "name", "region"]); w.writerows(sorted(out))
    merges = []
    for r in wb.worksheets[1].iter_rows(values_only=True):
        if r[0] and isinstance(r[0], int) and r[3] and str(r[3]).startswith("fus"):
            into = int(re.match(r"fus (\d+)", str(r[3])).group(1))
            if into != r[0]:
                merges.append((r[0], r[1], str(r[2]).strip(), into))
    with open(ROOT / f"data/kk/mergers-{year}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n"); w.writerow(["bfs", "canton", "name", "into"]); w.writerows(sorted(merges))
    print(year, len(out), "Gemeinden in divided cantons,", len(merges), "mergers")
